from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

from benchmark_construction.artifact_generation import (
    validate_artifact_blueprint,
)
from benchmark_construction.impress_artifact_catalog import (
    adapt_source_presentation,
    create_native_impress_seed,
    load_impress_artifact_catalog,
    materialize_impress_artifact_catalog,
)
from benchmark_construction.impress_setup_profiles import (
    apply_impress_setup_profiles,
)
from benchmark_construction.presentation_artifacts import inspect_pptx
from benchmark_construction.models import SkillRecord, SourceTask
from benchmark_construction.reference_annotation import (
    assemble_osworld_annotation_config,
    validate_annotation_blueprint,
)
from benchmark_construction.reference_applications import (
    get_reference_application,
)
from benchmark_construction.reference_generation import SkillSample
from benchmark_construction.reference_packages import (
    build_reference_package_request,
    validate_reference_package_response,
)
from benchmark_construction.reviewer_guide_generation import (
    ReviewerGuideJob,
    build_reviewer_guide_request,
)
from benchmark_construction.skill_extraction import build_skill_request
from benchmark_construction.reference_generation import load_skill_pool
from benchmark_construction.skill_readiness import eligible_skill_ids
from benchmark_construction.schema import load_schema, validate_payload
from scripts.python import build_impress_deferred_artifacts as deferred_builder


APP = "libreoffice_impress"


def _skill(index: int) -> SkillRecord:
    return SkillRecord(
        skill_id=f"source.skill-{index:02d}",
        app=APP,
        name=f"Impress skill {index}",
        procedure=("Select an editable object.", "Apply the requested operation."),
        efficiency_tip="Use the relevant sidebar control.",
        source_task_id="source",
        source_action_ids=(index - 1,),
    )


def _artifact_spec() -> dict:
    theme = {
        "primary_color": "2C5F2D",
        "secondary_color": "97BC62",
        "accent_color": "F2B134",
        "background_color": "F7F7F2",
        "title_font": "Liberation Sans",
        "body_font": "Liberation Sans",
    }
    return {
        "artifact_type": "pptx",
        "domain": "Community programs",
        "presentation_title": "Community Program Review",
        "slide_size": "wide",
        "theme": theme,
        "slides": [
            {
                "slide_number": 1,
                "purpose": "Introduce the review",
                "layout": "title_content",
                "object_plan": [
                    {
                        "semantic_id": "review_title",
                        "type": "text",
                        "purpose": "Presentation title",
                        "editable": True,
                    }
                ],
            },
            {
                "slide_number": 2,
                "purpose": "Compare participation",
                "layout": "data_story",
                "object_plan": [
                    {
                        "semantic_id": "participation_chart",
                        "type": "chart",
                        "purpose": "Participation comparison",
                        "editable": True,
                    }
                ],
            },
        ],
        "initial_state": ["Two coherent slides exist."],
        "must_not_be_completed": ["The requested formatting remains incomplete."],
        "generation_notes": ["Use editable static objects."],
    }


def _style() -> dict:
    return {
        "fill_color": "FFFFFF",
        "line_color": "D9E0E7",
        "text_color": "1F2937",
        "font_face": "Liberation Sans",
        "font_size": 18,
        "bold": False,
        "italic": False,
        "underline": False,
        "align": "left",
        "valign": "mid",
    }


def _content(**overrides) -> dict:
    content = {
        "text": None,
        "shape_kind": None,
        "image_kind": None,
        "rows": [],
        "categories": [],
        "series": [],
    }
    content.update(overrides)
    return content


def _package() -> dict:
    return {
        "reference_task_id": "reference-task-impress-r01-001",
        "app": APP,
        "task_instruction": "Format the title and update the comparison chart.",
        "required_skill_ids": ["source.skill-01", "source.skill-02"],
        "artifact_spec": _artifact_spec(),
        "operator_guide": {
            "recommended_demonstration": [
                {
                    "skill_id": "source.skill-01",
                    "operation_intent": "Format the title.",
                    "visible_success_signal": "The title has the requested style.",
                    "efficiency_tip": "Select the title once.",
                },
                {
                    "skill_id": "source.skill-02",
                    "operation_intent": "Update the chart.",
                    "visible_success_signal": "The chart shows the requested result.",
                    "efficiency_tip": "Edit the existing chart.",
                },
            ],
            "allowed_variation": "Equivalent English-UI paths are allowed.",
            "recording_start_state": "The presentation is open.",
            "recording_end_state": "The two requested edits are visible.",
        },
        "expected_incidental_operations": [],
    }


def _blueprint() -> dict:
    spec = _artifact_spec()
    return {
        "presentation_title": spec["presentation_title"],
        "slide_size": spec["slide_size"],
        "theme": spec["theme"],
        "slides": [
            {
                "slide_number": 1,
                "purpose": "Introduce the review",
                "background_color": "F7F7F2",
                "objects": [
                    {
                        "semantic_id": "review_title",
                        "type": "text",
                        "x": 0.7,
                        "y": 0.6,
                        "w": 7.0,
                        "h": 0.8,
                        "content": _content(text="Community Program Review"),
                        "style": _style(),
                    }
                ],
            },
            {
                "slide_number": 2,
                "purpose": "Compare participation",
                "background_color": "FFFFFF",
                "objects": [
                    {
                        "semantic_id": "participation_chart",
                        "type": "chart",
                        "x": 0.8,
                        "y": 1.4,
                        "w": 7.0,
                        "h": 4.5,
                        "content": _content(
                            categories=["Arts", "Sports"],
                            series=[{"name": "Participants", "values": [42, 58]}],
                        ),
                        "style": _style(),
                    }
                ],
            },
        ],
        "manual_setup_steps": [],
    }


def test_impress_profile_routes_app_contracts() -> None:
    profile = get_reference_application(APP)
    assert profile.artifact_type == "pptx"
    assert profile.packet_artifact_filename == "initial_artifact.pptx"
    assert profile.package_schema == "reference-package-impress-candidate.schema.json"


def test_impress_skill_and_package_requests_use_impress_prompts() -> None:
    source = SourceTask(
        task_id="source",
        app=APP,
        instruction="Move an image.",
        single_steps=("`CLICK` image", "`DRAG_TO` right side"),
    )
    skill_request = build_skill_request(source)
    assert skill_request.prompt_name == "extract_impress_skills.v1"

    sample = SkillSample(app=APP, skills=(_skill(1), _skill(2)))
    package_request = build_reference_package_request(sample, [source])
    assert package_request.prompt_name == "generate_impress_reference_package.v2"
    assert (
        package_request.response_schema["properties"]["artifact_spec"]["properties"][
            "artifact_type"
        ]["const"]
        == "pptx"
    )


def test_impress_package_and_artifact_blueprint_validate() -> None:
    sample = SkillSample(app=APP, skills=(_skill(1), _skill(2)))
    response = {
        "decision": "candidate",
        **_package(),
        "rejection_reason": "",
    }
    response.pop("reference_task_id")
    response.pop("app")
    validate_reference_package_response(sample, response)
    validate_artifact_blueprint(_package(), _blueprint())


def test_impress_blueprint_schema_allows_thirteen_editable_objects() -> None:
    blueprint = _blueprint()
    template = blueprint["slides"][0]["objects"][0]
    blueprint["slides"][0]["objects"] = [
        {
            **template,
            "semantic_id": f"editable_object_{index}",
            "x": 0.1 * index,
        }
        for index in range(13)
    ]

    validate_payload(
        blueprint,
        load_schema("impress-artifact-blueprint.schema.json"),
    )


def test_impress_candidate_schema_is_openai_strict_compatible() -> None:
    schema = load_schema("reference-package-impress-candidate.schema.json")
    assert set(schema["required"]) == set(schema["properties"])
    assert "artifact_strategy_ids" not in schema["properties"]


def test_impress_annotation_config_uses_pptx(tmp_path: Path) -> None:
    artifact = tmp_path / "initial_artifact.pptx"
    artifact.write_bytes(b"PK\x03\x04synthetic")
    digest = hashlib.sha256(artifact.read_bytes()).hexdigest()
    package = _package()
    manifest = {
        "reference_task_id": package["reference_task_id"],
        "artifact_path": artifact.name,
        "artifact_sha256": digest,
        "manual_setup_required": False,
        "manual_setup_steps": [],
    }
    blueprint = {
        "snapshot": APP,
        "artifact_slot": "initial_artifact",
        "guest_filename": "community_review.pptx",
        "open_after_upload": True,
        "ready_state_checks": ["Confirm the two-slide presentation is open."],
        "setup_notes": [],
    }
    validate_annotation_blueprint(blueprint, expected_app=APP)
    config = assemble_osworld_annotation_config(
        package, manifest, blueprint, repo_root=tmp_path
    )
    assert config["snapshot"] == APP
    assert config["related_apps"] == [APP]
    assert config["config"][0]["parameters"]["files"][0]["path"].endswith(".pptx")


def test_impress_annotation_config_supports_inputs_outputs_and_setup(
    tmp_path: Path,
) -> None:
    artifact = tmp_path / "initial_artifact.pptx"
    attachment = tmp_path / "source.png"
    artifact.write_bytes(b"PK\x03\x04synthetic")
    attachment.write_bytes(b"\x89PNG\r\n\x1a\nimage")
    manifest = {
        "reference_task_id": _package()["reference_task_id"],
        "artifact_path": artifact.name,
        "artifact_sha256": hashlib.sha256(artifact.read_bytes()).hexdigest(),
        "supporting_files": [
            {
                "role": "supporting_input",
                "local_path": attachment.name,
                "guest_path": "/home/user/Desktop/source.png",
                "sha256": hashlib.sha256(attachment.read_bytes()).hexdigest(),
            }
        ],
        "expected_outputs": [
            {
                "role": "png_export",
                "guest_path": "/home/user/Desktop/res.png",
                "local_filename": "res.png",
                "required": True,
            }
        ],
        "setup_profiles": ["hide_slides_pane"],
        "manual_setup_required": False,
        "manual_setup_steps": [],
    }
    blueprint = {
        "snapshot": APP,
        "artifact_slot": "initial_artifact",
        "guest_filename": "community_review.pptx",
        "open_after_upload": True,
        "ready_state_checks": ["The presentation is open."],
        "setup_notes": [],
    }

    config = assemble_osworld_annotation_config(
        _package(), manifest, blueprint, repo_root=tmp_path
    )

    assert len(config["config"][0]["parameters"]["files"]) == 2
    assert config["config"][1]["parameters"]["path"].endswith("community_review.pptx")
    annotation = config["reference_annotation"]
    assert annotation["supporting_files"][0]["guest_path"].endswith("source.png")
    assert annotation["expected_outputs"][0]["local_filename"] == "res.png"
    assert annotation["setup_profiles"] == ["hide_slides_pane"]


def test_deferred_catalog_covers_all_v2_readiness_overrides() -> None:
    root = Path("evaluation_examples/expert_skill_learning/impress_full_v1")
    skills = json.loads((root / "generated_v2/skill_pool.json").read_text())["skills"]
    policy = json.loads((root / "generated_v2/skill_readiness_policy.json").read_text())
    catalog = load_impress_artifact_catalog(
        root / "deferred_skill_artifact_catalog.json",
        schema_path=Path(
            "evaluation_examples/expert_skill_learning/schemas/"
            "impress-deferred-artifact-catalog.schema.json"
        ),
        known_skill_ids=[item["skill_id"] for item in skills],
        required_skill_ids=[item["skill_id"] for item in policy["overrides"]],
    )
    assert len(catalog["entries"]) == 18
    assert sum(len(item["skill_ids"]) for item in catalog["entries"]) == 21
    plan = json.loads((root / "deferred_skill_generation_plan.json").read_text())
    planned = {skill_id for group in plan["groups"] for skill_id in group["skill_ids"]}
    deferred = {item["skill_id"] for item in policy["overrides"]}
    assert deferred <= planned
    assert len(plan["groups"]) <= 20


def test_native_impress_seed_profiles_are_structurally_inspectable(
    tmp_path: Path,
) -> None:
    notes = tmp_path / "notes.pptx"
    portrait = tmp_path / "portrait.pptx"
    placeholders = tmp_path / "placeholders.pptx"
    create_native_impress_seed(notes, "notes_ready")
    create_native_impress_seed(portrait, "portrait_ready")
    create_native_impress_seed(placeholders, "title_body_placeholders")

    notes_inspection = inspect_pptx(notes)
    portrait_inspection = inspect_pptx(portrait)
    placeholder_inspection = inspect_pptx(placeholders)
    assert notes_inspection["notes_parts"]
    assert notes_inspection["slide_placeholder_count"] >= 2
    assert placeholder_inspection["empty_placeholder_count"] >= 1
    assert portrait_inspection["slide_count"] == 2
    assert portrait_inspection["orientation"] == "portrait"


def test_materializer_builds_native_strategy_without_network(tmp_path: Path) -> None:
    catalog = {
        "source_manifest": "source_tasks.json",
        "allowed_download_hosts": ["huggingface.co"],
        "entries": [
            {
                "strategy_id": "native-test",
                "skill_ids": ["skill-1"],
                "source_task_id": "source-1",
                "acquisition": "create_native",
                "native_profile": "generic_ui",
                "include_source_attachments": False,
                "setup_profiles": [],
                "expected_outputs": [],
                "adaptation_requirements": ["Open the presentation."],
                "validation_requirements": ["It opens."],
                "source_license": "Apache-2.0",
            }
        ],
    }
    manifest = materialize_impress_artifact_catalog(catalog, {}, tmp_path)
    primary = manifest["strategies"][0]["assets"][0]
    assert primary["validation"]["pptx"]["slide_count"] == 2
    assert Path(primary["local_path"]).is_file()


def test_hide_slides_pane_setup_is_allowlisted() -> None:
    calls = []
    env = SimpleNamespace(
        controller=SimpleNamespace(
            execute_python_command=lambda command: calls.append(command)
            or {"output": "setup-complete"}
        )
    )
    assert apply_impress_setup_profiles(env, ["hide_slides_pane"]) == [
        "hide_slides_pane"
    ]
    assert "presses=11" in calls[0]


def test_source_presentation_adaptation_scrubs_text_notes_and_raster_media(
    tmp_path: Path,
) -> None:
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches

    source = tmp_path / "source.pptx"
    image_path = tmp_path / "source-image.png"
    Image.new("RGB", (120, 80), "red").save(image_path)
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Sensitive source title"
    slide.shapes.add_picture(str(image_path), Inches(1), Inches(1))
    slide.notes_slide.notes_text_frame.text = "Sensitive private note"
    presentation.save(source)

    target = tmp_path / "adapted.pptx"
    adapt_source_presentation(source, target, "source_seed")
    inspection = inspect_pptx(target)
    assert "Sensitive source title" not in inspection["text"]
    assert inspection["media_parts"]
    assert inspection["slides"][0]["picture_count"] == 1
    with zipfile.ZipFile(source) as source_archive, zipfile.ZipFile(
        target
    ) as target_archive:
        assert source_archive.read("ppt/media/image1.png") != target_archive.read(
            "ppt/media/image1.png"
        )
        notes = target_archive.read("ppt/notesSlides/notesSlide1.xml")
        assert b"Sensitive private note" not in notes


def test_deferred_builder_emits_multi_input_and_output_manifest(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "source.pptx"
    support = tmp_path / "photo.png"
    create_native_impress_seed(source, "generic_ui")
    support.write_bytes(b"\x89PNG\r\n\x1a\nphoto")
    package = {
        **_package(),
        "artifact_strategy_ids": ["image-strategy", "export-strategy"],
    }
    materialized = {
        "strategies": [
            {
                "strategy_id": "image-strategy",
                "source_task_id": "source-image",
                "source_license": "Apache-2.0",
                "assets": [
                    {"role": "primary", "local_path": str(source)},
                    {
                        "role": "supporting",
                        "local_path": str(support),
                        "guest_path": "/home/user/Desktop/photo.png",
                    },
                ],
                "expected_outputs": [],
                "setup_profiles": [],
            },
            {
                "strategy_id": "export-strategy",
                "source_task_id": "source-export",
                "source_license": "Apache-2.0",
                "assets": [{"role": "primary", "local_path": str(source)}],
                "expected_outputs": [
                    {
                        "role": "png_export",
                        "guest_path": "/home/user/Desktop/res.png",
                        "local_filename": "res.png",
                        "required": True,
                    }
                ],
                "setup_profiles": ["hide_slides_pane"],
            },
        ]
    }

    def fake_render(_pptx, preview_dir, _seed, qa_path, **_kwargs):
        preview_dir.mkdir(parents=True)
        (preview_dir / "slide-01.png").write_bytes(b"png")
        qa_path.write_text("{}")

    monkeypatch.setattr(deferred_builder, "render_and_qa_presentation", fake_render)
    manifest = deferred_builder.build_deferred_artifacts(
        [package],
        materialized,
        tmp_path / "built",
        soffice="soffice",
        pdftoppm="pdftoppm",
    )
    entry = manifest["artifacts"][0]
    assert entry["supporting_files"][0]["guest_path"].endswith("photo.png")
    assert entry["expected_outputs"][0]["local_filename"] == "res.png"
    assert entry["setup_profiles"] == ["hide_slides_pane"]


def test_impress_reviewer_guide_request_uses_impress_prompt(tmp_path: Path) -> None:
    detail = tmp_path / "TASK_DETAIL.md"
    detail.write_text("# Detail\n\nImpress task detail.\n")
    job = ReviewerGuideJob(
        reference_task_id="reference-task-impress-r01-001",
        required_skill_ids=("source.skill-01", "source.skill-02"),
        task_detail_path=detail,
        app=APP,
    )
    request = build_reviewer_guide_request(job, detail.read_text())
    assert request.prompt_name == "generate_impress_reviewer_guide.v3"


def test_impress_pilot_readiness_keeps_deferred_skills_unselected() -> None:
    root = Path("evaluation_examples/expert_skill_learning/impress_full_v1")
    skills = load_skill_pool(root / "generated/skill_pool.json", expected_app=APP)
    eligible = eligible_skill_ids(
        root / "skill_readiness_policy.json",
        skills,
        expected_app=APP,
        statuses=["pilot_static"],
    )
    assert len(skills) == 66
    assert len(eligible) == 46
    assert "c59742c0-4323-4b9d-8a02-723c251deaa0.skill-01" not in eligible
    assert "04578141-1d42-4146-b9cf-6fab4ce5fd74.skill-01" in eligible
