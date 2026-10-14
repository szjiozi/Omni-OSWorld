"""Acquire and validate trusted artifacts for deferred Impress skills."""

from __future__ import annotations

import hashlib
import io
import json
import re
import urllib.request
import zipfile
from collections.abc import Callable, Iterable
from pathlib import Path, PurePosixPath
from typing import Any
from urllib.parse import urlparse
from xml.etree import ElementTree

from .presentation_artifacts import inspect_pptx
from .schema import validate_payload


SAFE_FILENAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._ -]{0,159}$")
RELATIONSHIP = (
    "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
)


def _sha256_bytes(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_impress_artifact_catalog(
    path: Path,
    *,
    schema_path: Path,
    known_skill_ids: Iterable[str] | None = None,
    required_skill_ids: Iterable[str] | None = None,
) -> dict[str, Any]:
    document = json.loads(path.read_text(encoding="utf-8"))
    schema = json.loads(schema_path.read_text(encoding="utf-8"))
    validate_payload(document, schema)

    strategy_ids: set[str] = set()
    catalog_skill_ids: set[str] = set()
    known = set(known_skill_ids or [])
    for entry in document["entries"]:
        strategy_id = entry["strategy_id"]
        if strategy_id in strategy_ids:
            raise ValueError(f"Duplicate Impress artifact strategy: {strategy_id}")
        strategy_ids.add(strategy_id)
        for skill_id in entry["skill_ids"]:
            if skill_id in catalog_skill_ids:
                raise ValueError(
                    f"Skill appears in multiple artifact strategies: {skill_id}"
                )
            if known and skill_id not in known:
                raise ValueError(
                    f"Artifact strategy references unknown skill: {skill_id}"
                )
            catalog_skill_ids.add(skill_id)

    required = set(required_skill_ids or [])
    missing = sorted(required - catalog_skill_ids)
    extra = sorted(catalog_skill_ids - required) if required else []
    if missing or extra:
        raise ValueError(
            "Deferred-skill catalog coverage mismatch: "
            f"missing={missing}, extra={extra}"
        )
    return document


def load_frozen_source_tasks(
    manifest_path: Path, *, repo_root: Path
) -> dict[str, dict[str, Any]]:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    source_root = (repo_root / manifest["vendored_source_root"]).resolve()
    tasks: dict[str, dict[str, Any]] = {}
    for summary in manifest["tasks"]:
        task_id = summary["task_id"]
        source_path = (source_root / summary["source_file"]).resolve()
        try:
            source_path.relative_to(source_root)
        except ValueError as exc:
            raise ValueError(
                f"Source task path escapes frozen source root: {source_path}"
            ) from exc
        if _sha256_file(source_path) != summary["source_sha256"]:
            raise ValueError(f"Frozen source task SHA256 mismatch: {task_id}")
        source_task = json.loads(source_path.read_text(encoding="utf-8"))
        if source_task.get("id") != task_id:
            raise ValueError(f"Frozen source task ID mismatch: {task_id}")
        tasks[task_id] = source_task
    return tasks


def source_downloads(source_task: dict[str, Any]) -> list[dict[str, str]]:
    downloads: list[dict[str, str]] = []
    for step in source_task.get("config", []):
        if step.get("type") != "download":
            continue
        for item in step.get("parameters", {}).get("files", []):
            if isinstance(item.get("url"), str) and isinstance(item.get("path"), str):
                downloads.append({"url": item["url"], "guest_path": item["path"]})
    return downloads


def _validate_download(
    url: str, guest_path: str, allowed_hosts: set[str]
) -> tuple[str, str]:
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname not in allowed_hosts:
        raise ValueError(f"Untrusted Impress artifact URL: {url}")
    guest = PurePosixPath(guest_path)
    if not guest.is_absolute() and guest.parts[:1] in {("Desktop",), ("Downloads",)}:
        guest = PurePosixPath("/home/user") / guest
    if guest.parent not in {
        PurePosixPath("/home/user/Desktop"),
        PurePosixPath("/home/user/Downloads"),
    }:
        raise ValueError(f"Unsupported guest artifact path: {guest_path}")
    filename = guest.name
    if not SAFE_FILENAME.fullmatch(filename):
        raise ValueError(f"Unsafe downloaded artifact filename: {filename!r}")
    return filename, str(guest)


def _default_downloader(url: str, allowed_hosts: set[str] | None = None) -> bytes:
    request = urllib.request.Request(
        url, headers={"User-Agent": "OSWorld-reference-builder/1.0"}
    )
    with urllib.request.urlopen(request, timeout=60) as response:
        if (
            allowed_hosts is not None
            and urlparse(response.geturl()).hostname not in allowed_hosts
        ):
            raise ValueError(
                f"Artifact download redirected to another host: {response.geturl()}"
            )
        return response.read()


def inspect_pptx_security(path: Path) -> dict[str, Any]:
    if not zipfile.is_zipfile(path):
        raise ValueError(f"PPTX is not a ZIP package: {path}")
    external_relationships: list[dict[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        names = archive.namelist()
        macro_parts = sorted(
            name for name in names if name.casefold().endswith("vbaproject.bin")
        )
        for name in names:
            if not name.endswith(".rels"):
                continue
            root = ElementTree.fromstring(archive.read(name))
            for node in root.iter(RELATIONSHIP):
                if node.attrib.get("TargetMode") == "External":
                    external_relationships.append(
                        {
                            "relationship_part": name,
                            "target": node.attrib.get("Target", ""),
                        }
                    )
    return {
        "macro_parts": macro_parts,
        "external_relationships": external_relationships,
    }


def create_native_impress_seed(path: Path, profile: str) -> None:
    """Create a small OOXML seed whose native structure supports a deferred workflow."""

    from pptx import Presentation
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt

    presentation = Presentation()
    if profile == "portrait_ready":
        presentation.slide_width, presentation.slide_height = (
            presentation.slide_height,
            presentation.slide_width,
        )
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[0])
    title_slide.shapes.title.text = "Quarterly Briefing"
    title_slide.placeholders[1].text = "Prepared for the review team"
    content_slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    content_slide.shapes.title.text = "Agenda"
    body = content_slide.placeholders[1].text_frame
    body.text = "Current status"
    body.add_paragraph().text = "Next steps"

    if profile == "title_body_placeholders":
        body.clear()
    elif profile == "notes_ready":
        title_slide.notes_slide.notes_text_frame.text = (
            "Introduce the briefing and confirm the audience."
        )
    elif profile == "summary_ready":
        for title in ("Results", "Risks", "Recommendations"):
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = title
            slide.placeholders[1].text_frame.text = f"Key points for {title.lower()}"
    elif profile == "slide_number_ready":
        text_box = content_slide.shapes.add_textbox(
            presentation.slide_width - Inches(1.0),
            presentation.slide_height - Inches(0.55),
            Inches(0.6),
            Inches(0.3),
        )
        paragraph = text_box.text_frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.RIGHT
        run = paragraph.add_run()
        run.text = "2"
        run.font.size = Pt(12)

    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(path)


def _neutralize_raster_media(name: str, content: bytes, index: int) -> bytes:
    """Replace source imagery with a text-free neutral visual of the same format."""

    suffix = PurePosixPath(name).suffix.casefold()
    if suffix not in {".gif", ".jpeg", ".jpg", ".png"}:
        return content
    from PIL import Image, ImageDraw

    try:
        with Image.open(io.BytesIO(content)) as source:
            width, height = source.size
    except OSError:
        return content
    width = max(32, min(width, 1920))
    height = max(32, min(height, 1080))
    palettes = (
        ("#F4F8FB", "#C9DDEA", "#F5D7A9"),
        ("#FAF7F3", "#D9CFDF", "#C7DFDF"),
        ("#F5F8F2", "#CEDFCE", "#EACDBB"),
    )
    background, primary, accent = palettes[index % len(palettes)]
    image = Image.new("RGB", (width, height), background)
    draw = ImageDraw.Draw(image)
    draw.polygon(
        [
            (0, height),
            (0, round(height * 0.58)),
            (round(width * 0.36), height),
        ],
        fill=primary,
    )
    draw.ellipse(
        (
            round(width * 0.76),
            round(height * 0.05),
            round(width * 1.04),
            round(height * 0.42),
        ),
        fill=accent,
    )
    draw.rectangle(
        (
            0,
            0,
            max(2, round(width * 0.025)),
            height,
        ),
        fill="#FFFFFF",
    )
    output = io.BytesIO()
    image_format = {
        ".gif": "GIF",
        ".jpeg": "JPEG",
        ".jpg": "JPEG",
        ".png": "PNG",
    }[suffix]
    image.save(output, format=image_format, quality=88)
    return output.getvalue()


def adapt_source_presentation(source: Path, target: Path, native_profile: str) -> None:
    """Create a privacy-safe text variant while preserving native OOXML parts."""

    synthetic = (
        ("Community Learning Showcase", "Program overview", "Schedule", "Resources"),
        ("Workshop Plan", "Opening session", "Hands-on activity", "Discussion"),
        ("Participation Snapshot", "Registrations", "Facilitators", "Next milestone"),
        ("Resource Guide", "Venue", "Materials", "Support contacts"),
        ("Closing Notes", "Key takeaway", "Follow-up", "Thank you"),
    )
    temp = target.with_suffix(".tmp.pptx")
    target.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(source) as input_archive:
        infos = input_archive.infolist()
        parts = {info.filename: input_archive.read(info.filename) for info in infos}

    media_index = 0
    for name, content in list(parts.items()):
        if not name.startswith("ppt/media/"):
            continue
        parts[name] = _neutralize_raster_media(name, content, media_index)
        media_index += 1

    removed_by_source_part: dict[str, set[str]] = {}
    for name, content in list(parts.items()):
        if not name.endswith(".rels"):
            continue
        root = ElementTree.fromstring(content)
        removed_ids = {
            node.attrib.get("Id", "")
            for node in root.iter(RELATIONSHIP)
            if node.attrib.get("TargetMode") == "External"
        }
        if not removed_ids:
            continue
        for node in list(root):
            if node.tag == RELATIONSHIP and node.attrib.get("Id") in removed_ids:
                root.remove(node)
        rel_path = PurePosixPath(name)
        source_part = str(rel_path.parent.parent / rel_path.name.removesuffix(".rels"))
        removed_by_source_part[source_part] = removed_ids
        parts[name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)

    for source_part, removed_ids in removed_by_source_part.items():
        if source_part not in parts:
            continue
        root = ElementTree.fromstring(parts[source_part])
        for parent in root.iter():
            for child in list(parent):
                if any(value in removed_ids for value in child.attrib.values()):
                    parent.remove(child)
        parts[source_part] = ElementTree.tostring(
            root, encoding="utf-8", xml_declaration=True
        )

    for name, content in list(parts.items()):
        if not re.fullmatch(r"ppt/slides/slide[1-9][0-9]*\.xml", name):
            continue
        match = re.search(r"slide([1-9][0-9]*)\.xml", name)
        slide_index = int(match.group(1)) - 1
        root = ElementTree.fromstring(content)
        replacements = synthetic[slide_index % len(synthetic)]
        text_index = 0
        for node in root.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        ):
            if (node.text or "").strip():
                node.text = (
                    replacements[text_index] if text_index < len(replacements) else ""
                )
                text_index += 1
        parts[name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)

    for name, content in list(parts.items()):
        if not re.fullmatch(r"ppt/notesSlides/notesSlide[1-9][0-9]*\.xml", name):
            continue
        root = ElementTree.fromstring(content)
        for node in root.iter(
            "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
        ):
            node.text = ""
        parts[name] = ElementTree.tostring(root, encoding="utf-8", xml_declaration=True)

    if "docProps/core.xml" in parts:
        root = ElementTree.fromstring(parts["docProps/core.xml"])
        safe_core_values = {
            "creator": "OSWorld",
            "lastModifiedBy": "OSWorld",
            "title": "Adapted Impress seed",
            "subject": "",
            "description": "",
            "keywords": "",
        }
        for node in root.iter():
            local_name = node.tag.rsplit("}", 1)[-1]
            if local_name in safe_core_values:
                node.text = safe_core_values[local_name]
        parts["docProps/core.xml"] = ElementTree.tostring(
            root, encoding="utf-8", xml_declaration=True
        )

    with zipfile.ZipFile(temp, "w", compression=zipfile.ZIP_DEFLATED) as output_archive:
        for info in infos:
            output_archive.writestr(info, parts[info.filename])
    temp.replace(target)

    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    from pptx.oxml.ns import qn

    presentation = Presentation(target)
    while len(presentation.slides) > 6:
        # python-pptx exposes no public slide-deletion API.
        slide_id = presentation.slides._sldIdLst[-1]
        presentation.part.drop_rel(slide_id.rId)
        del presentation.slides._sldIdLst[-1]
    if native_profile == "slide_number_ready":
        for slide in presentation.slides:
            for shape in list(slide.shapes):
                if not any(
                    field.attrib.get("type", "").casefold() == "slidenum"
                    for field in shape._element.iter(qn("a:fld"))
                ):
                    continue
                shape._element.getparent().remove(shape._element)
    if native_profile == "source_seed":
        for slide_index, slide in enumerate(presentation.slides):
            for shape in slide.shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    if shape.left < 0:
                        shape.width += shape.left
                        shape.left = 0
                    if shape.top < 0:
                        shape.height += shape.top
                        shape.top = 0
                    if shape.left + shape.width > presentation.slide_width:
                        shape.width = presentation.slide_width - shape.left
                    if shape.top + shape.height > presentation.slide_height:
                        shape.height = presentation.slide_height - shape.top
                if not shape.has_text_frame:
                    continue
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x33)
            has_text = any(
                shape.has_text_frame and shape.text.strip() for shape in slide.shapes
            )
            has_picture = any(
                shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes
            )
            if has_text or has_picture:
                continue
            replacements = synthetic[slide_index % len(synthetic)]
            empty_text_placeholders = [
                shape
                for shape in slide.placeholders
                if shape.has_text_frame and not shape.text.strip()
            ]
            for shape, replacement in zip(empty_text_placeholders, replacements):
                shape.text = replacement
    if native_profile == "title_body_placeholders":
        for slide in presentation.slides:
            content_placeholders = [
                shape
                for shape in slide.placeholders
                if shape.has_text_frame
                and shape.placeholder_format.type
                in {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT}
            ]
            if content_placeholders:
                content_placeholders[0].text_frame.clear()
                break
    if native_profile == "notes_ready":
        for slide in presentation.slides:
            slide.notes_slide.notes_text_frame.clear()
    presentation.save(target)


def _validate_materialized_asset(
    path: Path, *, allow_external_relationships: bool
) -> dict[str, Any]:
    suffix = path.suffix.casefold()
    content = path.read_bytes()
    if not content:
        raise ValueError(f"Materialized artifact is empty: {path}")
    if suffix == ".pptx":
        security = inspect_pptx_security(path)
        if security["macro_parts"]:
            raise ValueError(f"PPTX contains macros: {security['macro_parts']}")
        if security["external_relationships"] and not allow_external_relationships:
            raise ValueError(
                "PPTX contains external relationships: "
                f"{security['external_relationships']}"
            )
        return {"pptx": inspect_pptx(path), "security": security}
    if suffix == ".png" and not content.startswith(b"\x89PNG\r\n\x1a\n"):
        raise ValueError(f"Downloaded PNG has an invalid signature: {path}")
    if suffix == ".mp3" and not (
        content.startswith(b"ID3")
        or content[:2] in {b"\xff\xfb", b"\xff\xf3", b"\xff\xf2"}
    ):
        raise ValueError(f"Downloaded MP3 has an invalid signature: {path}")
    return {"size_bytes": len(content)}


def _validate_native_profile(profile: str, inspection: dict[str, Any]) -> None:
    if profile == "title_body_placeholders" and (
        inspection["slide_placeholder_count"] < 2
        or not inspection["empty_placeholder_count"]
    ):
        raise ValueError(
            "Native profile requires title/body placeholders and one empty content slot"
        )
    if profile == "notes_ready" and not inspection["notes_parts"]:
        raise ValueError("Native notes profile has no notes-slide parts")
    if profile == "master_ready" and not inspection["master_parts"]:
        raise ValueError("Native master profile has no slide master")
    if (
        profile == "slide_number_ready"
        and not inspection["master_layout_slide_number_field_count"]
    ):
        raise ValueError(
            "Native slide-number profile has no slide-number field in a master or layout"
        )
    if profile == "summary_ready" and inspection["slide_count"] < 3:
        raise ValueError("Summary Slide profile requires at least three source slides")


def materialize_impress_artifact_catalog(
    catalog: dict[str, Any],
    source_tasks: dict[str, dict[str, Any]],
    output_dir: Path,
    *,
    downloader: Callable[[str], bytes] = _default_downloader,
    allow_external_relationships: bool = False,
) -> dict[str, Any]:
    output_dir.mkdir(parents=True, exist_ok=True)
    allowed_hosts = set(catalog["allowed_download_hosts"])
    materialized: list[dict[str, Any]] = []
    for entry in catalog["entries"]:
        strategy_dir = output_dir / entry["strategy_id"]
        strategy_dir.mkdir(parents=True, exist_ok=True)
        assets: list[dict[str, Any]] = []
        if entry["acquisition"] == "create_native":
            guest_path = "/home/user/Desktop/initial_artifact.pptx"
            target = strategy_dir / "initial_artifact.pptx"
            create_native_impress_seed(target, entry["native_profile"])
            assets.append(
                {
                    "role": "primary",
                    "guest_path": guest_path,
                    "local_path": str(target),
                    "source_url": None,
                }
            )
        else:
            source_task = source_tasks.get(entry["source_task_id"])
            if source_task is None:
                raise ValueError(f"Unknown source task: {entry['source_task_id']}")
            downloads = source_downloads(source_task)
            pptx_downloads = [
                item
                for item in downloads
                if item["guest_path"].casefold().endswith(".pptx")
            ]
            if not pptx_downloads:
                raise ValueError(
                    f"Source task has no PPTX input: {entry['source_task_id']}"
                )
            primary_url = pptx_downloads[0]["url"]
            selected = (
                downloads if entry["include_source_attachments"] else pptx_downloads[:1]
            )
            seen_names: set[str] = set()
            downloaded: list[dict[str, Any]] = []
            for item in selected:
                filename, normalized_guest_path = _validate_download(
                    item["url"], item["guest_path"], allowed_hosts
                )
                if filename.casefold() in seen_names:
                    raise ValueError(f"Duplicate source asset filename: {filename}")
                seen_names.add(filename.casefold())
                target = (
                    strategy_dir / "_source" / filename
                    if item["url"] == primary_url
                    else strategy_dir / filename
                )
                target.parent.mkdir(parents=True, exist_ok=True)
                if target.is_file():
                    content = target.read_bytes()
                elif downloader is _default_downloader:
                    content = _default_downloader(item["url"], allowed_hosts)
                else:
                    content = downloader(item["url"])
                if not target.is_file():
                    target.write_bytes(content)
                downloaded.append(
                    {
                        "role": "primary"
                        if item["url"] == primary_url
                        else "supporting",
                        "guest_path": normalized_guest_path,
                        "local_path": str(target),
                        "source_url": item["url"],
                        "source_sha256": _sha256_bytes(content),
                    }
                )
            source_primary = [item for item in downloaded if item["role"] == "primary"][
                0
            ]
            adapted_path = strategy_dir / "initial_artifact.pptx"
            adapt_source_presentation(
                Path(source_primary["local_path"]),
                adapted_path,
                entry["native_profile"],
            )
            assets.append(
                {
                    **source_primary,
                    "guest_path": "/home/user/Desktop/initial_artifact.pptx",
                    "local_path": str(adapted_path),
                    "derived_from": source_primary["local_path"],
                }
            )
            assets.extend(item for item in downloaded if item["role"] == "supporting")

        for asset in assets:
            target = Path(asset["local_path"])
            asset["sha256"] = _sha256_file(target)
            asset["validation"] = _validate_materialized_asset(
                target,
                allow_external_relationships=allow_external_relationships,
            )
        primary_validation = next(
            asset["validation"]["pptx"]
            for asset in assets
            if asset["role"] == "primary"
        )
        _validate_native_profile(entry["native_profile"], primary_validation)
        if sum(asset["role"] == "primary" for asset in assets) != 1:
            raise ValueError(
                f"Strategy must materialize one primary PPTX: {entry['strategy_id']}"
            )
        materialized.append(
            {
                "strategy_id": entry["strategy_id"],
                "skill_ids": entry["skill_ids"],
                "source_task_id": entry["source_task_id"],
                "source_license": entry["source_license"],
                "native_profile": entry["native_profile"],
                "setup_profiles": entry["setup_profiles"],
                "expected_outputs": entry["expected_outputs"],
                "adaptation_requirements": entry["adaptation_requirements"],
                "validation_requirements": entry["validation_requirements"],
                "assets": assets,
            }
        )
    manifest = {
        "schema_version": "1.0",
        "app": "libreoffice_impress",
        "catalog_source": catalog["source_manifest"],
        "strategies": materialized,
    }
    (output_dir / "materialized_artifacts.json").write_text(
        json.dumps(manifest, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    return manifest
