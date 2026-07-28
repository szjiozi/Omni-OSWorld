#!/usr/bin/env python3
"""Validate the four Phase 0 tasks locally and optionally on Daytona."""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import cv2


REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

TASK_ROOT = REPO_ROOT / "evaluation_examples" / "video_learning"
DEFAULT_REPORT = REPO_ROOT / "results" / "phase0_task_validation.json"
FIXTURE_DIR = TASK_ROOT / "fixtures" / "generated"


def _absolute(path: str) -> Path:
    candidate = Path(path)
    return candidate if candidate.is_absolute() else REPO_ROOT / candidate


def _load_tasks() -> list[tuple[Path, dict[str, Any]]]:
    paths = sorted((TASK_ROOT / "examples").glob("*/*.json"))
    return [
        (path, json.loads(path.read_text(encoding="utf-8")))
        for path in paths
    ]


def _input_artifact(task: dict[str, Any]) -> Path:
    for step in task["config"]:
        if step["type"] != "upload_file":
            continue
        for file_config in step["parameters"]["files"]:
            suffix = Path(file_config["local_path"]).suffix.lower()
            if suffix in {".pptx", ".xlsx"}:
                return _absolute(file_config["local_path"])
    raise ValueError(f"Task {task['id']} has no Office input artifact")


def _expected_artifact(task: dict[str, Any]) -> Path:
    expected = task["evaluator"]["expected"]
    if expected["type"] != "local_file":
        raise ValueError(f"Task {task['id']} expected artifact is not local")
    return _absolute(expected["path"])


def _phase0_compare_table(
    result_path: Path,
    expected_path: Path,
    options: dict[str, Any],
) -> float:
    from openpyxl import load_workbook

    result = load_workbook(result_path)
    expected = load_workbook(expected_path)
    if result.sheetnames != expected.sheetnames:
        return 0.0

    def cell_style(cell, prop):
        if prop == "number_format":
            return cell.number_format if cell.value is not None else None
        if prop == "font_bold":
            return cell.font.bold if cell.value is not None else None
        if prop == "font_color":
            color = cell.font.color
            return str(color.rgb) if cell.value is not None and color else None
        if prop == "bgcolor":
            return (
                str(cell.fill.fgColor.rgb)
                if cell.value is not None and cell.fill
                else None
            )
        raise ValueError(f"Unsupported Phase 0 style property: {prop}")

    for rule in options["rules"]:
        result_sheet = result[result.sheetnames[int(rule["sheet_idx0"])]]
        expected_index = str(rule["sheet_idx1"])
        expected_sheet = expected[
            expected.sheetnames[int(expected_index.removeprefix("EI"))]
        ]
        if rule["type"] == "sheet_data":
            result_values = [
                [cell.value for cell in row]
                for row in result_sheet.iter_rows()
            ]
            expected_values = [
                [cell.value for cell in row]
                for row in expected_sheet.iter_rows()
            ]
            if result_values != expected_values:
                return 0.0
        elif rule["type"] == "style":
            for result_row, expected_row in zip(
                result_sheet.iter_rows(),
                expected_sheet.iter_rows(),
            ):
                for result_cell, expected_cell in zip(result_row, expected_row):
                    for prop in rule["props"]:
                        if cell_style(result_cell, prop) != cell_style(
                            expected_cell,
                            prop,
                        ):
                            return 0.0
        else:
            raise ValueError(
                f"Unsupported Phase 0 table rule: {rule['type']}"
            )
    return 1.0


def _phase0_compare_pptx(result_path: Path, expected_path: Path) -> float:
    from pptx import Presentation

    result = Presentation(result_path)
    expected = Presentation(expected_path)
    if len(result.slides) != len(expected.slides):
        return 0.0

    def background_rgb(slide):
        fill = slide.background.fill
        if fill.type is None:
            return None
        try:
            return str(fill.fore_color.rgb)
        except (AttributeError, TypeError):
            return None

    def text_signature(shape):
        if not getattr(shape, "has_text_frame", False):
            return None
        paragraphs = []
        for paragraph in shape.text_frame.paragraphs:
            runs = []
            for run in paragraph.runs:
                color = run.font.color
                try:
                    color_rgb = str(color.rgb)
                except (AttributeError, TypeError):
                    color_rgb = None
                runs.append(
                    (
                        run.text,
                        run.font.name,
                        run.font.size.pt if run.font.size else None,
                        run.font.bold,
                        run.font.italic,
                        color_rgb,
                    )
                )
            paragraphs.append((paragraph.text, tuple(runs)))
        return tuple(paragraphs)

    for result_slide, expected_slide in zip(result.slides, expected.slides):
        if background_rgb(result_slide) != background_rgb(expected_slide):
            return 0.0
        if len(result_slide.shapes) != len(expected_slide.shapes):
            return 0.0
        for result_shape, expected_shape in zip(
            result_slide.shapes,
            expected_slide.shapes,
        ):
            result_geometry = (
                result_shape.shape_type,
                result_shape.left,
                result_shape.top,
                result_shape.width,
                result_shape.height,
            )
            expected_geometry = (
                expected_shape.shape_type,
                expected_shape.left,
                expected_shape.top,
                expected_shape.width,
                expected_shape.height,
            )
            if result_geometry != expected_geometry:
                return 0.0
            if text_signature(result_shape) != text_signature(expected_shape):
                return 0.0
    return 1.0


def _metric_scores(task: dict[str, Any]) -> tuple[float, float, str]:
    initial = _input_artifact(task)
    expected = _expected_artifact(task)
    evaluator = task["evaluator"]
    options = evaluator.get("options", {})
    if evaluator["func"] == "compare_pptx_files":
        try:
            from desktop_env.evaluators.metrics.slides import compare_pptx_files
        except ModuleNotFoundError as exc:
            if exc.name not in {"formulas", "rapidfuzz"}:
                raise
            return (
                _phase0_compare_pptx(initial, expected),
                _phase0_compare_pptx(expected, expected),
                "phase0_compatible",
            )
        else:
            return (
                float(compare_pptx_files(str(initial), str(expected), **options)),
                float(compare_pptx_files(str(expected), str(expected), **options)),
                "osworld",
            )
    if evaluator["func"] == "compare_table":
        try:
            from desktop_env.evaluators.metrics.table import compare_table
        except ModuleNotFoundError as exc:
            if exc.name not in {"formulas", "rapidfuzz"}:
                raise
            return (
                _phase0_compare_table(initial, expected, options),
                _phase0_compare_table(expected, expected, options),
                "phase0_compatible",
            )
        else:
            return (
                float(compare_table(str(initial), str(expected), **options)),
                float(compare_table(str(expected), str(expected), **options)),
                "osworld",
            )
    raise ValueError(f"Unsupported Phase 0 metric: {evaluator['func']}")


def _video_metadata(path: Path) -> dict[str, Any]:
    capture = cv2.VideoCapture(str(path))
    try:
        if not capture.isOpened():
            raise ValueError(f"Video is not decodable: {path}")
        frame_count = int(capture.get(cv2.CAP_PROP_FRAME_COUNT))
        fps = float(capture.get(cv2.CAP_PROP_FPS))
        if frame_count <= 0 or fps <= 0:
            raise ValueError(f"Video has invalid timing metadata: {path}")
        return {
            "frames": frame_count,
            "fps": fps,
            "duration_seconds": frame_count / fps,
        }
    finally:
        capture.release()


def _validate_fixture_manifest() -> dict[str, Any]:
    manifest_path = FIXTURE_DIR / "fixture_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    for name, expected in manifest["files"].items():
        path = FIXTURE_DIR / name
        if not path.is_file():
            raise FileNotFoundError(f"Fixture manifest entry is missing: {path}")
        actual_hash = hashlib.sha256(path.read_bytes()).hexdigest()
        if actual_hash != expected["sha256"]:
            raise ValueError(f"Fixture hash mismatch: {path}")
        if path.stat().st_size != expected["bytes"]:
            raise ValueError(f"Fixture size mismatch: {path}")
    return {
        "path": str(manifest_path.relative_to(REPO_ROOT)),
        "files": len(manifest["files"]),
        "verified": True,
    }


def _normalize_local_paths(task: dict[str, Any]) -> dict[str, Any]:
    task = copy.deepcopy(task)
    for step in task["config"]:
        if step["type"] != "upload_file":
            continue
        for file_config in step["parameters"]["files"]:
            file_config["local_path"] = str(_absolute(file_config["local_path"]))
    expected = task["evaluator"].get("expected")
    if expected and expected.get("type") == "local_file":
        expected["path"] = str(_absolute(expected["path"]))
    return task


def _daytona_validate(
    tasks: list[tuple[Path, dict[str, Any]]],
) -> list[dict[str, Any]]:
    from desktop_env.desktop_env import DesktopEnv

    env = DesktopEnv(
        provider_name="daytona",
        os_type="Ubuntu",
        action_space="pyautogui",
        headless=True,
        require_a11y_tree=False,
    )
    results = []
    try:
        for path, raw_task in tasks:
            task = _normalize_local_paths(raw_task)
            env.reset(task_config=task)
            initial_score = float(env.evaluate())

            gold_task = copy.deepcopy(task)
            input_path = str(_input_artifact(raw_task))
            expected_path = str(_expected_artifact(raw_task))
            for step in gold_task["config"]:
                if step["type"] != "upload_file":
                    continue
                for file_config in step["parameters"]["files"]:
                    if file_config["local_path"] == input_path:
                        file_config["local_path"] = expected_path
            env.reset(task_config=gold_task)
            gold_score = float(env.evaluate())
            if initial_score >= 1.0:
                raise AssertionError(f"{path}: initial artifact already passes")
            if gold_score < 1.0:
                raise AssertionError(f"{path}: gold artifact scored {gold_score}")
            results.append(
                {
                    "task_id": task["id"],
                    "initial_score": initial_score,
                    "gold_score": gold_score,
                }
            )
    finally:
        env.close()
    return results


def validate(*, generate: bool, daytona: bool) -> dict[str, Any]:
    if generate:
        from scripts.python.generate_phase0_fixtures import generate as generate_fixtures

        generate_fixtures()

    from desktop_env.trajectory import validate_document

    tasks = _load_tasks()
    if len(tasks) != 4:
        raise AssertionError(f"Expected exactly 4 Phase 0 tasks, found {len(tasks)}")

    report: dict[str, Any] = {
        "schema_version": "1.0",
        "validated_at": datetime.now(timezone.utc).isoformat(),
        "status": "running",
        "fixture_manifest": _validate_fixture_manifest(),
        "tasks": [],
    }
    for path, task in tasks:
        validate_document(task, "task")
        skill_path = _absolute(task["video_learning"]["skill_ir"])
        skill = json.loads(skill_path.read_text(encoding="utf-8"))
        validate_document(skill, "skill")

        source_artifact = _absolute(
            task["video_learning"]["source_artifact"]
        )
        demo_result_artifact = _absolute(
            task["video_learning"]["demo_result_artifact"]
        )
        if not source_artifact.is_file():
            raise FileNotFoundError(source_artifact)
        if not demo_result_artifact.is_file():
            raise FileNotFoundError(demo_result_artifact)

        demo_path = _absolute(task["video_learning"]["demo_video"]["path"])
        video = _video_metadata(demo_path)
        initial_score, gold_score, metric_engine = _metric_scores(task)
        if initial_score >= 1.0:
            raise AssertionError(f"{path}: initial artifact already passes")
        if gold_score < 1.0:
            raise AssertionError(f"{path}: gold artifact scored {gold_score}")
        report["tasks"].append(
            {
                "task_id": task["id"],
                "path": str(path.relative_to(REPO_ROOT)),
                "initial_score": initial_score,
                "gold_score": gold_score,
                "metric_engine": metric_engine,
                "video": video,
                "skill": str(skill_path.relative_to(REPO_ROOT)),
            }
        )

    if daytona:
        report["daytona"] = _daytona_validate(tasks)
    report["status"] = "passed"
    return report


def _write_report(path: Path, report: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp_path = path.with_suffix(path.suffix + ".tmp")
    tmp_path.write_text(
        json.dumps(report, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    os.replace(tmp_path, path)


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--generate",
        action="store_true",
        help="Regenerate deterministic fixtures before validation.",
    )
    parser.add_argument(
        "--daytona",
        action="store_true",
        help="Also run full task setup/evaluator validation on Daytona.",
    )
    parser.add_argument(
        "--report",
        type=Path,
        default=DEFAULT_REPORT,
    )
    return parser.parse_args()


def main() -> int:
    args = _parse_args()
    try:
        report = validate(generate=args.generate, daytona=args.daytona)
    except Exception as exc:
        report = {
            "schema_version": "1.0",
            "validated_at": datetime.now(timezone.utc).isoformat(),
            "status": "failed",
            "error": f"{type(exc).__name__}: {exc}",
        }
        _write_report(args.report, report)
        print(f"[FAIL] {exc}", file=sys.stderr)
        return 1

    _write_report(args.report, report)
    print(f"[PASS] validated {len(report['tasks'])} Phase 0 tasks")
    print(f"Report: {args.report}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
